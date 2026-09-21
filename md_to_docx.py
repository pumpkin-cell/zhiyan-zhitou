# md_to_docx.py —— 把商业计划书.md 转成正式 Word 文档（含封面页 + 目录页，去掉 Markdown 标识符）
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import sys

SRC = sys.argv[1] if len(sys.argv) > 1 else "商业计划书.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "商业计划书.docx"


def set_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_rich_text(paragraph, text, size=11, bold_all=False):
    """处理 **加粗** 与 `代码`，去掉 Markdown 标识符后写入段落。"""
    for i, seg in enumerate(text.split('**')):
        if seg == '':
            continue
        seg = seg.replace('`', '')  # 去反引号
        run = paragraph.add_run(seg)
        set_font(run, size=size, bold=(bold_all or i % 2 == 1))


def parse_table(lines):
    """把连续 | 行解析成 (表头, 数据行)。"""
    header = [c.strip() for c in lines[0].strip().strip('|').split('|')]
    rows = []
    for ln in lines[2:]:  # 跳过表头与分隔线
        rows.append([c.strip() for c in ln.strip().strip('|').split('|')])
    return header, rows


def add_cover(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("智研智投"), name="黑体", size=36, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("—— 面向金融量化投研工作全流程的智能体系统"), name="黑体", size=22, bold=True)
    doc.add_paragraph()
    for line in ["中国国际大学生创新大赛（2026）", "产教协同创新组 · 新工科", "命题企业：达观数据有限公司"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(line), size=14)
    for _ in range(4):
        doc.add_paragraph()
    for line in ["团队名称：____________________", "团队成员：____________________",
                 "指导老师：____________________", "2026 年 9 月"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(line), size=12)
    doc.add_page_break()


def setup_styles(doc):
    """按用户格式设置样式：正文宋体11pt/1.15行距，标题宋体14/13pt粗体黑色。"""
    normal = doc.styles['Normal']
    normal.font.name = '宋体'
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.15
    for sname, sz in [('Heading 1', 14), ('Heading 2', 13)]:
        h = doc.styles[sname]
        h.font.name = '宋体'
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def add_toc_field(doc):
    """插入 TOC 自动目录（带页码），Word 打开后按 F9 更新。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("目  录"), name="黑体", size=18, bold=True)
    p2 = doc.add_paragraph()
    r = p2.add_run()
    fld = OxmlElement('w:fldChar')
    fld.set(qn('w:fldCharType'), 'begin')
    fld.set(qn('w:dirty'), 'true')
    r._r.append(fld)
    r2 = p2.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    r2._r.append(instr)
    r3 = p2.add_run()
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    r3._r.append(sep)
    r4 = p2.add_run("（打开 Word 后 Ctrl+A 全选，再按 F9 更新目录页码）")
    set_font(r4, size=10)
    r5 = p2.add_run()
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    r5._r.append(end)
    doc.add_page_break()


def convert():
    doc = Document()
    lines = open(SRC, encoding='utf-8').read().split('\n')

    setup_styles(doc)
    add_cover(doc)
    add_toc_field(doc)

    i = 0
    n = len(lines)
    heading_seen = False
    while i < n:
        ln = lines[i]
        if ln.startswith('# '):
            pass  # 主标题已在封面
        elif ln.startswith('## '):
            if heading_seen:
                doc.add_page_break()
            heading_seen = True
            doc.add_heading(ln[3:].strip(), level=1)
        elif ln.startswith('### '):
            doc.add_heading(ln[4:].strip(), level=2)
        elif ln.startswith('|'):
            tbl = [ln]
            j = i + 1
            while j < n and lines[j].startswith('|'):
                tbl.append(lines[j])
                j += 1
            header, rows = parse_table(tbl)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = 'Table Grid'
            for c, h in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ''
                add_rich_text(cell.paragraphs[0], h, size=10, bold_all=True)
            for row in rows:
                cells = table.add_row().cells
                for c, val in enumerate(row):
                    if c < len(header):
                        cells[c].text = ''
                        add_rich_text(cells[c].paragraphs[0], val, size=10)
            i = j - 1
        elif ln.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_rich_text(p, ln[2:].strip())
        elif ln.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_rich_text(p, ln[2:].strip())
        elif ln.strip() in ('---', '--- '):
            pass
        elif ln.strip() == '':
            pass
        else:
            if ln.strip():
                p = doc.add_paragraph()
                add_rich_text(p, ln.strip())
        i += 1

    doc.save(OUT)
    print(f"✅ 已生成 Word 文档: {OUT}（封面 + 目录 + 正文，Markdown 标识符已去除）")


if __name__ == '__main__':
    convert()
